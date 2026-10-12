"""Structure-aware PDF / DOCX / text extraction for advanced RAG."""

from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
import re
from typing import Any, Dict, List, Optional

ALLOWED_SUFFIXES = {".pdf", ".docx", ".txt", ".md"}
MAX_UPLOAD_BYTES = 25 * 1024 * 1024  # 25 MB


class DocumentParseError(Exception):
    """Raised when an upload cannot be parsed into usable text."""

    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.status_code = status_code


@dataclass
class DocumentSection:
    """Represents a structural block or section in a document."""

    content: str
    page_number: Optional[int] = None
    section_title: Optional[str] = None
    element_type: str = "paragraph"  # "heading", "paragraph", "table", "header"
    metadata: Dict[str, Any] = field(default_factory=dict)

    @property
    def token_estimate(self) -> int:
        return max(1, len(self.content) // 4)

    def to_citation(self) -> str:
        parts = []
        if self.page_number is not None:
            parts.append(f"Page {self.page_number}")
        if self.section_title:
            parts.append(f"Section: {self.section_title}")
        return ", ".join(parts) if parts else "Document"


@dataclass
class ParsedDocument:
    """Rich document container holding text, page breakdown, and sections."""

    filename: str
    total_pages: int
    raw_text: str
    sections: List[DocumentSection] = field(default_factory=list)
    page_texts: Dict[int, str] = field(default_factory=dict)
    metadata: Dict[str, Any] = field(default_factory=dict)

    @property
    def total_words(self) -> int:
        return len(self.raw_text.split())

    @property
    def total_characters(self) -> int:
        return len(self.raw_text)

    def get_full_text_with_citations(self) -> str:
        """Returns document text with explicit structural markers for LLM ingestion."""
        blocks: List[str] = []
        for sec in self.sections:
            header = f"[{sec.to_citation()}]"
            blocks.append(f"{header}\n{sec.content}")
        return "\n\n".join(blocks)


def _suffix(filename: str | None) -> str:
    if not filename:
        raise DocumentParseError("Uploaded file is missing a filename.")
    return Path(filename).suffix.lower()


def parse_document_to_structured(filename: str, data: bytes) -> ParsedDocument:
    """Parses PDF or DOCX file into a structured ParsedDocument retaining pages and headings."""
    suffix = _suffix(filename)
    if suffix not in ALLOWED_SUFFIXES:
        raise DocumentParseError(
            f"Unsupported file type '{suffix or 'unknown'}'. Use .pdf, .docx, .txt, or .md."
        )
    if not data:
        raise DocumentParseError("The uploaded file is empty.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise DocumentParseError(f"File exceeds the {MAX_UPLOAD_BYTES // (1024 * 1024)} MB upload limit.")

    if suffix == ".pdf":
        return _parse_pdf_structured(filename, data)
    elif suffix == ".docx":
        return _parse_docx_structured(filename, data)
    else:
        text = data.decode("utf-8", errors="replace")
        return parse_text_to_structured(filename, text)


def parse_text_to_structured(title: str, text: str) -> ParsedDocument:
    """Converts raw or markdown text into structured sections."""
    lines = text.splitlines()
    sections: List[DocumentSection] = []
    current_title: Optional[str] = None
    current_lines: List[str] = []

    heading_pattern = re.compile(r"^(#{1,6}\s+|[0-9]+(\.[0-9]+)*\s+[A-Z]|[A-Z\s]{4,}:)")

    def flush_section():
        nonlocal current_lines, current_title
        content = " ".join(" ".join(current_lines).split()).strip()
        if content:
            sections.append(
                DocumentSection(
                    content=content,
                    page_number=1,
                    section_title=current_title,
                    element_type="paragraph",
                )
            )
        current_lines = []

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        if heading_pattern.match(stripped) and len(stripped) < 120:
            flush_section()
            current_title = stripped.lstrip("#").strip()
            sections.append(
                DocumentSection(
                    content=current_title,
                    page_number=1,
                    section_title=current_title,
                    element_type="heading",
                )
            )
        else:
            current_lines.append(stripped)

    flush_section()

    if not sections and text.strip():
        sections.append(
            DocumentSection(
                content=" ".join(text.split()).strip(),
                page_number=1,
                section_title=None,
                element_type="paragraph",
            )
        )

    clean_raw = "\n\n".join(s.content for s in sections if s.content.strip())
    return ParsedDocument(
        filename=title,
        total_pages=1,
        raw_text=clean_raw,
        sections=sections,
        page_texts={1: clean_raw},
    )


def extract_text_from_bytes(filename: str, data: bytes) -> str:
    """Backward-compatible helper returning cleaned raw string."""
    doc = parse_document_to_structured(filename, data)
    cleaned = " ".join(doc.raw_text.split()).strip()
    if not cleaned:
        raise DocumentParseError("No extractable text found in the uploaded file.")
    return cleaned


def _parse_pdf_structured(filename: str, data: bytes) -> ParsedDocument:
    from pypdf import PdfReader

    try:
        reader = PdfReader(BytesIO(data))
    except Exception as exc:
        raise DocumentParseError(f"Unable to read PDF: {exc}") from exc

    sections: List[DocumentSection] = []
    page_texts: Dict[int, str] = {}
    total_pages = len(reader.pages)

    current_section_title: Optional[str] = None
    heading_pattern = re.compile(r"^([0-9]+(\.[0-9]+)*\s+[A-Z][A-Za-z0-9\s]{2,60}|[A-Z\s]{4,40})$")

    for idx, page in enumerate(reader.pages, start=1):
        raw_page_text = page.extract_text() or ""
        cleaned_page = " ".join(raw_page_text.split()).strip()
        page_texts[idx] = cleaned_page

        # Break page into rough paragraphs
        lines = [line.strip() for line in raw_page_text.splitlines() if line.strip()]
        current_para: List[str] = []

        for line in lines:
            if heading_pattern.match(line) and len(line) < 80:
                if current_para:
                    sections.append(
                        DocumentSection(
                            content=" ".join(current_para),
                            page_number=idx,
                            section_title=current_section_title,
                            element_type="paragraph",
                        )
                    )
                    current_para = []
                current_section_title = line
                sections.append(
                    DocumentSection(
                        content=line,
                        page_number=idx,
                        section_title=line,
                        element_type="heading",
                    )
                )
            else:
                current_para.append(line)

        if current_para:
            sections.append(
                DocumentSection(
                    content=" ".join(current_para),
                    page_number=idx,
                    section_title=current_section_title,
                    element_type="paragraph",
                )
            )

    all_raw = "\n\n".join(sec.content for sec in sections if sec.content.strip())
    if not all_raw.strip():
        raise DocumentParseError("No readable text found in the PDF.")

    return ParsedDocument(
        filename=filename,
        total_pages=total_pages,
        raw_text=all_raw,
        sections=sections,
        page_texts=page_texts,
    )


def _parse_docx_structured(filename: str, data: bytes) -> ParsedDocument:
    from docx import Document

    try:
        document = Document(BytesIO(data))
    except Exception as exc:
        raise DocumentParseError(f"Unable to read DOCX: {exc}") from exc

    sections: List[DocumentSection] = []
    current_section_title: Optional[str] = None

    for p in document.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        style_name = (p.style.name if p.style else "").lower()
        is_heading = "heading" in style_name or "title" in style_name

        if is_heading:
            current_section_title = txt
            sections.append(
                DocumentSection(
                    content=txt,
                    page_number=None,
                    section_title=txt,
                    element_type="heading",
                )
            )
        else:
            sections.append(
                DocumentSection(
                    content=txt,
                    page_number=None,
                    section_title=current_section_title,
                    element_type="paragraph",
                )
            )

    # Tables
    for table in document.tables:
        rows_content = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                rows_content.append(" | ".join(cells))
        if rows_content:
            table_text = "\n".join(rows_content)
            sections.append(
                DocumentSection(
                    content=table_text,
                    page_number=None,
                    section_title=current_section_title,
                    element_type="table",
                )
            )

    all_raw = "\n\n".join(sec.content for sec in sections if sec.content.strip())
    if not all_raw.strip():
        raise DocumentParseError("No readable text found in the DOCX.")

    return ParsedDocument(
        filename=filename,
        total_pages=1,
        raw_text=all_raw,
        sections=sections,
        page_texts={1: all_raw},
    )


def combine_reference(reference_text: str | None, file_text: str | None) -> str:
    """Combines explicit input text and extracted file text safely."""
    chunks = [c.strip() for c in (reference_text or "", file_text or "") if c and c.strip()]
    return "\n\n".join(chunks).strip()
